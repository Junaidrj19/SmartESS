"""Build the AgenticX SmartESS technical deck.

Visual structure follows the six-slide BurningGuard / SIH reference
(title, proposed solution, technical approach, feasibility, impact, references).
Content is SmartESS only. No SIH branding. Claims are limited to what the
repository actually implements and to verified corpus sources.
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

SW, SH = 20.0, 11.25
FONT = "Calibri"

NAVY = RGBColor(0x00, 0x35, 0x7B)
INK = RGBColor(0x1B, 0x2A, 0x41)
BLUE = RGBColor(0x0A, 0x3D, 0x8F)
ROYAL = RGBColor(0x15, 0x5A, 0xD6)
TITLE = RGBColor(0x0A, 0x24, 0x66)
RED = RGBColor(0xC0, 0x39, 0x2B)
ORANGE = RGBColor(0xE6, 0x7E, 0x22)
GREEN = RGBColor(0x1B, 0x8A, 0x4A)
TEAL = RGBColor(0x0E, 0x74, 0x86)
PURPLE = RGBColor(0x6C, 0x3B, 0xB5)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
MUTED = RGBColor(0x4A, 0x5A, 0x6A)
PALE = RGBColor(0xE7, 0xF3, 0xFC)
PALE_GREEN = RGBColor(0xE7, 0xF7, 0xEE)
PALE_ORANGE = RGBColor(0xFF, 0xF4, 0xE5)
PALE_RED = RGBColor(0xFD, 0xEE, 0xEC)
PALE_PURPLE = RGBColor(0xF3, 0xEB, 0xFB)
PALE_TEAL = RGBColor(0xE5, 0xF6, 0xF7)
CARD = RGBColor(0xF5, 0xFB, 0xFE)
LINE = RGBColor(0xC5, 0xD6, 0xE6)
SOFT = RGBColor(0xD6, 0xE6, 0xF5)


def _anchor(tf, anchor: str) -> None:
    body = tf._txBody.find(qn("a:bodyPr"))
    body.set("anchor", anchor)


def rrect(slide, x, y, w, h, fill, line=None, radius=0.12, weight=1.0):
    sh = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(weight)
    try:
        sh.adjustments[0] = radius
    except Exception:
        pass
    return sh


def oval(slide, x, y, w, h, fill, line=None):
    sh = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(1.25)
    return sh


def rect(slide, x, y, w, h, fill, line=None):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(1)
    return sh


def chevron(slide, x, y, w, h, fill):
    sh = slide.shapes.add_shape(MSO_SHAPE.CHEVRON, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.fill.background()
    return sh


def tb(slide, x, y, w, h, paragraphs, align=PP_ALIGN.LEFT, anchor="t"):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_left = Inches(0.06)
    tf.margin_right = Inches(0.05)
    tf.margin_top = Inches(0.03)
    tf.margin_bottom = Inches(0.02)
    _anchor(tf, anchor)
    for i, para in enumerate(paragraphs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = para.get("align", align)
        p.space_before = Pt(para.get("before", 0))
        p.space_after = Pt(para.get("after", 1))
        runs = para.get("runs")
        if runs is None:
            runs = [para]
        for spec in runs:
            run = p.add_run()
            run.text = spec["text"]
            run.font.size = Pt(spec.get("size", 13))
            run.font.bold = spec.get("bold", False)
            run.font.italic = spec.get("italic", False)
            run.font.color.rgb = spec.get("color", INK)
            run.font.name = FONT
            if spec.get("url"):
                run.hyperlink.address = spec["url"]
    return box


def bg(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def team_badge(slide, x=0.32, y=0.22):
    oval(slide, x, y, 1.55, 0.72, WHITE, RGBColor(0x6C, 0x5C, 0xE7))
    tb(
        slide,
        x,
        y + 0.08,
        1.55,
        0.56,
        [
            {"text": "Gradient", "size": 12, "bold": True, "color": PURPLE, "align": PP_ALIGN.CENTER, "after": 0},
            {"text": "Descents", "size": 12, "bold": True, "color": PURPLE, "align": PP_ALIGN.CENTER, "before": 0},
        ],
        align=PP_ALIGN.CENTER,
        anchor="ctr",
    )


def agenticx_mark(slide, x=17.55, y=0.18):
    rrect(slide, x, y, 2.15, 0.78, NAVY, radius=0.18)
    tb(
        slide,
        x,
        y + 0.08,
        2.15,
        0.64,
        [
            {"text": "AGENTICX", "size": 16, "bold": True, "color": WHITE, "align": PP_ALIGN.CENTER, "after": 0},
            {"text": "TECHNICAL PRESENTATION", "size": 8, "bold": True, "color": RGBColor(0xB9, 0xD4, 0xFF), "align": PP_ALIGN.CENTER},
        ],
        align=PP_ALIGN.CENTER,
        anchor="ctr",
    )


def slide_title(slide, title, subtitle, title_color=TITLE, y=0.18):
    tb(
        slide,
        2.1,
        y,
        15.3,
        0.48,
        [{"text": title, "size": 32, "bold": True, "color": title_color, "align": PP_ALIGN.CENTER, "after": 0}],
        align=PP_ALIGN.CENTER,
    )
    tb(
        slide,
        2.4,
        y + 0.46,
        14.7,
        0.32,
        [{"text": subtitle, "size": 13, "color": MUTED, "align": PP_ALIGN.CENTER, "after": 0}],
        align=PP_ALIGN.CENTER,
    )


def build():
    prs = Presentation()
    prs.slide_width = Inches(SW)
    prs.slide_height = Inches(SH)
    prs.core_properties.title = "SmartESS — AgenticX Technical Presentation"
    prs.core_properties.author = "Gradient Descents"
    prs.core_properties.subject = (
        "AI-assisted scientific investigation for electronic-component reliability"
    )
    prs.core_properties.category = "AgenticX"
    blank = prs.slide_layouts[6]
    slide_title_page(prs.slides.add_slide(blank))
    slide_solution(prs.slides.add_slide(blank))
    slide_technical(prs.slides.add_slide(blank))
    slide_feasibility(prs.slides.add_slide(blank))
    slide_impact(prs.slides.add_slide(blank))
    slide_references(prs.slides.add_slide(blank))
    return prs


def slide_title_page(slide):
    bg(slide, WHITE)
    tb(
        slide,
        0.6,
        0.28,
        14.5,
        0.7,
        [{"text": "AGENTICX", "size": 40, "bold": True, "color": TITLE, "align": PP_ALIGN.CENTER, "after": 0}],
        align=PP_ALIGN.CENTER,
    )
    agenticx_mark(slide, 17.35, 0.28)
    tb(
        slide,
        0.7,
        1.15,
        11.2,
        0.42,
        [{"text": "TECHNICAL PRESENTATION", "size": 22, "bold": True, "color": INK, "after": 0}],
    )
    # metadata
    rows = [
        ("Project", "SmartESS"),
        ("Statement", "AI-assisted scientific investigation for electronic-component reliability screening"),
        ("Focus", "SiC power modules — behavioral evidence, not specification-limit checking alone"),
        ("Domain", "Electronic component reliability  /  engineering intelligence"),
        ("Core", "Agentic AI  +  deterministic analysis  +  evidence-grounded investigation"),
        ("Decision", "Engineer review remains the final decision layer"),
        ("Team", "Gradient Descents"),
    ]
    y = 1.75
    for label, value in rows:
        oval(slide, 0.85, y + 0.08, 0.18, 0.18, ROYAL)
        tb(
            slide,
            1.2,
            y,
            10.6,
            0.72 if len(value) > 70 else 0.48,
            [
                {
                    "runs": [
                        {"text": label + "  —  ", "size": 18, "bold": True, "color": INK},
                        {"text": value, "size": 18, "bold": False, "color": INK},
                    ],
                    "after": 0,
                }
            ],
        )
        y += 0.78 if len(value) > 70 else 0.58

    # right emblem
    hexagon = slide.shapes.add_shape(
        MSO_SHAPE.HEXAGON, Inches(12.55), Inches(1.85), Inches(6.7), Inches(7.7)
    )
    hexagon.fill.solid()
    hexagon.fill.fore_color.rgb = RGBColor(0xEE, 0xF3, 0xF8)
    hexagon.line.fill.background()

    tb(
        slide,
        13.15,
        2.15,
        5.5,
        0.7,
        [
            {"text": "EPISTEMIC SEPARATION", "size": 14, "bold": True, "color": TITLE, "align": PP_ALIGN.CENTER, "after": 0},
            {"text": "Four kinds of statement stay distinct", "size": 12, "color": MUTED, "align": PP_ALIGN.CENTER},
        ],
        align=PP_ALIGN.CENTER,
    )
    layers = [
        (NAVY, "MEASURED DATA", "What the module record shows"),
        (TEAL, "DETERMINISTIC INFERENCE", "What analysis computed"),
        (PURPLE, "RETRIEVED EVIDENCE", "What curated sources support"),
        (ORANGE, "LLM HYPOTHESIS", "A bounded explanation — not a fact"),
        (GREEN, "ENGINEER DECISION", "Human judgment is final"),
    ]
    ly = 3.05
    for color, name, note in layers:
        rrect(slide, 13.35, ly, 5.1, 0.95, color, radius=0.15)
        tb(
            slide,
            13.45,
            ly + 0.08,
            4.9,
            0.8,
            [
                {"text": name, "size": 14, "bold": True, "color": WHITE, "align": PP_ALIGN.CENTER, "after": 0},
                {"text": note, "size": 12, "color": RGBColor(0xF4, 0xF8, 0xFF), "align": PP_ALIGN.CENTER},
            ],
            align=PP_ALIGN.CENTER,
            anchor="ctr",
        )
        ly += 1.08

    tb(
        slide,
        0.7,
        10.55,
        18.5,
        0.4,
        [
            {
                "text": "SmartESS   ·   Detect  →  Measure  →  Evidence  →  Hypothesis  →  Validate  →  Report  →  Human decision",
                "size": 14,
                "bold": True,
                "color": NAVY,
                "align": PP_ALIGN.LEFT,
            }
        ],
    )


def slide_solution(slide):
    bg(slide, RGBColor(0xF3, 0xFB, 0xFE))
    rect(slide, 0, 0, SW, 1.15, WHITE)
    rrect(slide, 0.28, 0.18, 3.55, 0.82, NAVY, radius=0.16)
    tb(
        slide,
        0.35,
        0.22,
        3.4,
        0.74,
        [
            {"text": "SmartESS", "size": 22, "bold": True, "color": WHITE, "align": PP_ALIGN.CENTER, "after": 0},
            {"text": "DETECT  ·  MEASURE  ·  INVESTIGATE  ·  REPORT", "size": 9, "bold": True, "color": RGBColor(0xC5, 0xDC, 0xFF), "align": PP_ALIGN.CENTER},
        ],
        align=PP_ALIGN.CENTER,
        anchor="ctr",
    )
    tb(
        slide,
        4.0,
        0.12,
        12.0,
        0.55,
        [{"text": "Proposed Solution", "size": 32, "bold": True, "color": TITLE, "align": PP_ALIGN.CENTER, "after": 0}],
        align=PP_ALIGN.CENTER,
    )
    tb(
        slide,
        4.0,
        0.64,
        12.0,
        0.38,
        [
            {
                "text": "Scientific investigation of SiC power-module reliability",
                "size": 14,
                "color": MUTED,
                "align": PP_ALIGN.CENTER,
            }
        ],
        align=PP_ALIGN.CENTER,
    )
    agenticx_mark(slide, 16.55, 0.2)

    # LEFT
    rrect(slide, 0.22, 1.32, 5.15, 8.05, WHITE, LINE, 0.08, 1.0)
    rect(slide, 0.22, 1.32, 0.1, 8.05, ROYAL)
    tb(
        slide,
        0.4,
        1.4,
        4.85,
        0.36,
        [{"text": "OUR PROPOSED SOLUTION", "size": 14, "bold": True, "color": BLUE, "after": 0}],
    )
    tb(
        slide,
        0.38,
        1.78,
        4.85,
        3.55,
        [
            {
                "text": "A SiC module can stay inside conventional limits and still show an abnormal trajectory. Threshold checks do not explain that behavior.",
                "size": 12,
                "color": INK,
                "after": 6,
            },
            {
                "text": "Engineers correlate electrical and thermal signals, test context, and literature by hand. Reasoning is fragmented and hard to audit.",
                "size": 12,
                "color": INK,
                "after": 6,
            },
            {
                "text": "SmartESS combines deterministic analysis, anomaly and degradation evidence, a curated engineering knowledge base, and specialized agents. The result is an evidence-grounded hypothesis and a structured investigation report.",
                "size": 12,
                "color": INK,
                "after": 6,
            },
            {
                "text": "The language model does not diagnose hardware. The engineer reviews the report and decides.",
                "size": 12,
                "bold": True,
                "color": NAVY,
                "after": 2,
            },
        ],
    )
    # gap table
    rrect(slide, 0.4, 5.45, 4.8, 0.42, RGBColor(0xE8, 0xEE, 0xF6), radius=0.08)
    tb(
        slide,
        0.45,
        5.48,
        2.2,
        0.36,
        [{"text": "CURRENT GAP", "size": 11, "bold": True, "color": RED, "align": PP_ALIGN.CENTER}],
        align=PP_ALIGN.CENTER,
        anchor="ctr",
    )
    tb(
        slide,
        2.7,
        5.48,
        2.4,
        0.36,
        [{"text": "SMARTESS RESPONSE", "size": 11, "bold": True, "color": GREEN, "align": PP_ALIGN.CENTER}],
        align=PP_ALIGN.CENTER,
        anchor="ctr",
    )
    gaps = [
        ("Static limit checks", "Behavioral anomaly evidence"),
        ("In-spec drift unexplained", "Measured degradation indicators"),
        ("Sources not traceable", "Retrieval with provenance"),
        ("One opaque AI answer", "Separated data, model, evidence, hypothesis"),
    ]
    gy = 5.95
    for left, right in gaps:
        tb(slide, 0.42, gy, 2.25, 0.55, [{"text": left, "size": 11, "color": INK, "after": 0}])
        tb(slide, 2.65, gy, 0.3, 0.4, [{"text": "→", "size": 14, "bold": True, "color": ROYAL, "after": 0}])
        tb(slide, 2.95, gy, 2.2, 0.7, [{"text": right, "size": 11, "bold": True, "color": NAVY, "after": 0}])
        gy += 0.78

    # CENTER architecture
    rrect(slide, 5.52, 1.32, 8.55, 8.05, WHITE, SOFT, 0.06, 1.25)
    tb(
        slide,
        5.7,
        1.4,
        8.2,
        0.3,
        [{"text": "INVESTIGATION ARCHITECTURE", "size": 12, "bold": True, "color": BLUE, "align": PP_ALIGN.CENTER, "after": 0}],
        align=PP_ALIGN.CENTER,
    )
    stages = [
        ("01", "ENGINEERING DATA", "Module, test, and telemetry records", False),
        ("02", "DETERMINISTIC ANALYSIS", "Features, rules, anomaly score", False),
        ("03", "ANOMALY / DEGRADATION EVIDENCE", "Measured deviation, not a diagnosis", False),
        ("04", "AGENTIC INVESTIGATION", "Question and evidence requirements", True),
        ("05", "KNOWLEDGE RETRIEVAL", "Curated corpus, source provenance", True),
        ("06", "EVIDENCE SYNTHESIS", "Retrieved claims stay distinct from data", True),
        ("07", "HYPOTHESIS", "Bounded engineering explanation", True),
        ("08", "VALIDATION", "Support, contradiction, limitations", True),
        ("09", "TECHNICAL REPORT", "Observed · computed · retrieved · hypothesized", False),
        ("10", "ENGINEER DECISION", "Human judgment is the last layer", False),
    ]
    sy = 1.78
    for num, name, note, agentic in stages:
        fill = RGBColor(0xEE, 0xE7, 0xFB) if agentic else PALE
        accent = PURPLE if agentic else ROYAL
        rrect(slide, 5.85, sy, 7.9, 0.68, fill, radius=0.12)
        oval(slide, 6.0, sy + 0.12, 0.44, 0.44, accent)
        tb(
            slide,
            6.0,
            sy + 0.16,
            0.44,
            0.36,
            [{"text": num, "size": 10, "bold": True, "color": WHITE, "align": PP_ALIGN.CENTER, "after": 0}],
            align=PP_ALIGN.CENTER,
            anchor="ctr",
        )
        tb(
            slide,
            6.55,
            sy + 0.04,
            5.3,
            0.6,
            [
                {"text": name, "size": 12, "bold": True, "color": INK, "after": 0},
                {"text": note, "size": 11, "color": MUTED, "before": 0, "after": 0},
            ],
            anchor="ctr",
        )
        if agentic:
            tb(
                slide,
                11.7,
                sy + 0.14,
                1.85,
                0.4,
                [{"text": "AGENT", "size": 10, "bold": True, "color": PURPLE, "align": PP_ALIGN.RIGHT, "after": 0}],
                align=PP_ALIGN.RIGHT,
                anchor="ctr",
            )
        sy += 0.73

    # RIGHT innovations
    rrect(slide, 14.22, 1.32, 5.52, 8.05, RGBColor(0xE8, 0xF3, 0xFC), radius=0.08)
    rrect(slide, 14.4, 1.48, 5.16, 0.7, ROYAL, radius=0.12)
    tb(
        slide,
        14.5,
        1.5,
        4.95,
        0.66,
        [
            {"text": "Key Innovation", "size": 18, "bold": True, "color": WHITE, "align": PP_ALIGN.CENTER, "after": 0},
            {"text": "What makes SmartESS different", "size": 11, "color": RGBColor(0xD6, 0xE6, 0xFF), "align": PP_ALIGN.CENTER},
        ],
        align=PP_ALIGN.CENTER,
        anchor="ctr",
    )
    innovations = [
        ("01", "Behavioral anomaly detection", "Beyond specification-limit checking. In-spec abnormal trajectories are evidence."),
        ("02", "Evidence-grounded investigation", "Curated engineering knowledge with source provenance. Not open-web search."),
        ("03", "Multi-agent reasoning", "Investigation, evidence, hypothesis, and report agents — not one LLM call."),
        ("04", "Epistemic separation", "Observed data, computed results, retrieved sources, and hypotheses stay distinct."),
        ("05", "Auditable engineering reports", "Evidence → hypothesis → validation → recommendation, for engineer review."),
    ]
    iy = 2.35
    for num, title, body in innovations:
        rrect(slide, 14.4, iy, 5.16, 1.28, WHITE, radius=0.1)
        oval(slide, 14.55, iy + 0.38, 0.52, 0.52, NAVY)
        tb(
            slide,
            14.55,
            iy + 0.46,
            0.52,
            0.38,
            [{"text": num, "size": 11, "bold": True, "color": WHITE, "align": PP_ALIGN.CENTER, "after": 0}],
            align=PP_ALIGN.CENTER,
        )
        tb(
            slide,
            15.2,
            iy + 0.08,
            4.2,
            1.12,
            [
                {"text": title, "size": 13, "bold": True, "color": BLUE, "after": 1},
                {"text": body, "size": 12, "color": INK, "after": 0},
            ],
        )
        iy += 1.36

    # bottom strip
    rect(slide, 0, 9.5, SW, 1.75, NAVY)
    outcomes = [
        ("DETECT", "Identify abnormal behavior"),
        ("MEASURE", "Quantify the deviation"),
        ("INVESTIGATE", "Trace relevant evidence"),
        ("EXPLAIN", "Bounded hypotheses"),
        ("VALIDATE", "Support and contradictions"),
        ("REPORT", "Auditable investigation"),
    ]
    ox = 0.3
    for i, (name, note) in enumerate(outcomes):
        rrect(slide, ox, 9.72, 2.85, 1.05, RGBColor(0x0C, 0x4A, 0x96), radius=0.12)
        tb(
            slide,
            ox + 0.08,
            9.8,
            2.7,
            0.9,
            [
                {"text": name, "size": 14, "bold": True, "color": WHITE, "align": PP_ALIGN.CENTER, "after": 0},
                {"text": note, "size": 11, "color": RGBColor(0xD5, 0xE6, 0xFF), "align": PP_ALIGN.CENTER},
            ],
            align=PP_ALIGN.CENTER,
            anchor="ctr",
        )
        ox += 3.05
    # The center panel extends under the footer. Clip by drawing footer last — already done.
    # But center panel was drawn before footer and is y=1.32 h=8.05 → ends at 9.37. Good, footer starts 9.5.


def slide_technical(slide):
    bg(slide, RGBColor(0xF4, 0xF9, 0xFD))
    team_badge(slide)
    slide_title(
        slide,
        "TECHNICAL APPROACH",
        "FROM MODULE OBSERVATIONS TO AN AUDITABLE INVESTIGATION",
        title_color=RED,
        y=0.12,
    )
    agenticx_mark(slide)

    # technology bar
    rrect(slide, 0.25, 1.15, 19.5, 0.78, RGBColor(0xE3, 0xF1, 0xFB), radius=0.1)
    tb(
        slide,
        0.35,
        1.22,
        2.15,
        0.64,
        [
            {"text": "TECHNOLOGIES", "size": 11, "bold": True, "color": BLUE, "align": PP_ALIGN.CENTER, "after": 0},
            {"text": "In this prototype", "size": 10, "color": MUTED, "align": PP_ALIGN.CENTER},
        ],
        align=PP_ALIGN.CENTER,
        anchor="ctr",
    )
    techs = [
        ("Python", "Analysis runtime"),
        ("FastAPI", "Investigation API"),
        ("scikit-learn", "Isolation Forest"),
        ("Next.js", "Engineer UI"),
        ("LangGraph", "Agent workflow"),
        ("ChromaDB", "Vector retrieval"),
        ("MiniLM", "Embeddings"),
        ("OpenRouter", "Llama 3.3 70B"),
    ]
    tx = 2.55
    for name, sub in techs:
        rrect(slide, tx, 1.26, 2.1, 0.56, WHITE, radius=0.12)
        tb(
            slide,
            tx,
            1.28,
            2.1,
            0.52,
            [
                {"text": name, "size": 12, "bold": True, "color": NAVY, "align": PP_ALIGN.CENTER, "after": 0},
                {"text": sub, "size": 9, "color": MUTED, "align": PP_ALIGN.CENTER},
            ],
            align=PP_ALIGN.CENTER,
            anchor="ctr",
        )
        tx += 2.14

    columns = [
        (
            "1",
            ROYAL,
            RGBColor(0xE3, 0xF2, 0xFD),
            "ENGINEERING DATA",
            "Module measurements, test context, telemetry",
            [
                "Module profile, test profile, and time-aligned telemetry.",
                "Representative parameters: RDS(on), Vth, IGSS, IDSS, VDS(on), Tj, Tc, thermal resistance.",
                "Labeled synthetic trajectories exist for development. They are not treated as field failures.",
            ],
        ),
        (
            "2",
            GREEN,
            PALE_GREEN,
            "DATA & SIGNAL PROCESSING",
            "Prepare trajectories before any model call",
            [
                "Cleaning and data-quality gates: PASS / WARNING / BLOCKED.",
                "Normalization and versioned observation features.",
                "Trajectory construction and module-level comparison.",
                "Quality failures stop the investigation rather than being guessed through.",
            ],
        ),
        (
            "3",
            PURPLE,
            PALE_PURPLE,
            "DETERMINISTIC ANALYSIS",
            "Produces measurable evidence",
            [
                "Frozen Isolation Forest anomaly score on observation features.",
                "Deterministic tools: drift, slope, percent change, population comparison, temperature dependence, change point, correlation, acceptance limits, degradation rate.",
                "Output is a computed indicator. It is not a root cause.",
            ],
        ),
        (
            "4",
            TEAL,
            PALE_TEAL,
            "INVESTIGATION ORCHESTRATION",
            "Investigation Agent builds the case",
            [
                "Observation → investigation question → evidence requirements.",
                "Context comes from the trajectory, anomaly summary, and module/test profiles.",
                "This stage does not ask the language model to diagnose the hardware.",
            ],
        ),
        (
            "5",
            ORANGE,
            PALE_ORANGE,
            "EVIDENCE & KNOWLEDGE",
            "Retrieved text is not module data",
            [
                "Curated corpus: 19 verified documents in ChromaDB (360 chunks).",
                "Documents → extraction → chunking → metadata / provenance → embeddings → retriever.",
                "Each retrieved claim is tied to its source. Open-web search is not the knowledge base.",
            ],
        ),
        (
            "6",
            RED,
            PALE_RED,
            "HYPOTHESIS, VALIDATION, REPORT",
            "Specialized agents, then the engineer",
            [
                "Hypothesis Agent proposes candidate explanations from the evidence.",
                "Validation checks support, contradictions, confidence, and limitations.",
                "Report Agent writes a structured technical report.",
                "Engineer review is the decision. A hypothesis is never stored as a measurement.",
            ],
        ),
    ]
    cw = 3.15
    gap = 0.1
    x0 = 0.25
    for i, (num, color, fill, title, sub, bullets) in enumerate(columns):
        x = x0 + i * (cw + gap)
        rrect(slide, x, 2.08, cw, 7.55, fill, radius=0.08)
        oval(slide, x + cw / 2 - 0.28, 2.2, 0.56, 0.56, color)
        tb(
            slide,
            x + cw / 2 - 0.28,
            2.28,
            0.56,
            0.42,
            [{"text": num, "size": 16, "bold": True, "color": WHITE, "align": PP_ALIGN.CENTER, "after": 0}],
            align=PP_ALIGN.CENTER,
        )
        tb(
            slide,
            x + 0.08,
            2.84,
            cw - 0.16,
            0.85,
            [
                {"text": title, "size": 13, "bold": True, "color": color, "align": PP_ALIGN.CENTER, "after": 2},
                {"text": sub, "size": 11, "color": INK, "align": PP_ALIGN.CENTER, "after": 0},
            ],
            align=PP_ALIGN.CENTER,
        )
        by = 3.75
        for b in bullets:
            tb(
                slide,
                x + 0.12,
                by,
                cw - 0.22,
                1.35,
                [{"text": "•  " + b, "size": 11, "color": INK, "after": 0}],
            )
            by += 1.15 if len(b) > 90 else 0.95

    rect(slide, 0, 9.78, SW, 1.47, WHITE)
    tb(
        slide,
        0.3,
        9.84,
        19.4,
        0.28,
        [
            {
                "text": "DETECTION   →   MEASUREMENT   →   EVIDENCE   →   HYPOTHESIS   →   VALIDATION   →   REPORT   →   HUMAN DECISION",
                "size": 13,
                "bold": True,
                "color": NAVY,
                "align": PP_ALIGN.CENTER,
                "after": 0,
            }
        ],
        align=PP_ALIGN.CENTER,
    )
    kinds = [
        (NAVY, "FACT", "Observed record"),
        (TEAL, "MODEL OUTPUT", "Computed score / statistic"),
        (PURPLE, "RETRIEVED", "Source-backed passage"),
        (ORANGE, "HYPOTHESIS", "Proposed explanation"),
        (GREEN, "DECISION", "Engineer judgment"),
    ]
    kx = 0.4
    for color, name, note in kinds:
        rrect(slide, kx, 10.22, 3.75, 0.78, color, radius=0.12)
        tb(
            slide,
            kx + 0.08,
            10.28,
            3.6,
            0.66,
            [
                {"text": name, "size": 13, "bold": True, "color": WHITE, "align": PP_ALIGN.CENTER, "after": 0},
                {"text": note, "size": 11, "color": RGBColor(0xF3, 0xF7, 0xFF), "align": PP_ALIGN.CENTER},
            ],
            align=PP_ALIGN.CENTER,
            anchor="ctr",
        )
        kx += 3.9


def slide_feasibility(slide):
    bg(slide, RGBColor(0xF7, 0xFC, 0xFE))
    team_badge(slide)
    slide_title(
        slide,
        "FEASIBILITY & VIABILITY",
        "From a working investigation prototype to an engineering decision-support layer",
        y=0.12,
    )
    agenticx_mark(slide)

    # 01
    rrect(slide, 0.28, 1.22, 10.55, 3.55, PALE_GREEN, RGBColor(0xB7, 0xE0, 0xC8), 0.08, 1.25)
    oval(slide, 0.48, 1.4, 0.7, 0.7, GREEN)
    tb(
        slide,
        0.48,
        1.52,
        0.7,
        0.48,
        [{"text": "01", "size": 16, "bold": True, "color": WHITE, "align": PP_ALIGN.CENTER, "after": 0}],
        align=PP_ALIGN.CENTER,
    )
    tb(
        slide,
        1.35,
        1.42,
        6.5,
        0.55,
        [{"text": "FEASIBLE NOW", "size": 22, "bold": True, "color": GREEN, "after": 0}],
    )
    tb(
        slide,
        8.3,
        1.45,
        2.3,
        0.7,
        [{"text": "Implemented\nand demonstrated\nin this repository", "size": 11, "color": MUTED, "align": PP_ALIGN.RIGHT, "after": 0}],
        align=PP_ALIGN.RIGHT,
    )
    feasible = [
        "Versioned module, test, and telemetry contracts.",
        "Data-quality gates and versioned observation features.",
        "Frozen Isolation Forest anomaly scoring, plus deterministic engineering tools.",
        "LangGraph workflow: Investigation, Evidence, Hypothesis, and Report agents.",
        "Curated corpus in ChromaDB with provenance. 19 verified documents.",
        "Structured investigation report. FastAPI and a Next.js engineer workstation.",
        "Deterministic results and LLM hypotheses are stored as different kinds of statement.",
    ]
    fy = 2.15
    for item in feasible:
        tb(slide, 0.55, fy, 10.0, 0.34, [{"text": "✓   " + item, "size": 13, "color": INK, "after": 0}])
        fy += 0.35

    # 02
    rrect(slide, 11.0, 1.22, 8.72, 3.55, PALE_ORANGE, RGBColor(0xF5, 0xD0, 0xA8), 0.08, 1.25)
    oval(slide, 11.2, 1.4, 0.7, 0.7, ORANGE)
    tb(
        slide,
        11.2,
        1.52,
        0.7,
        0.48,
        [{"text": "02", "size": 16, "bold": True, "color": WHITE, "align": PP_ALIGN.CENTER, "after": 0}],
        align=PP_ALIGN.CENTER,
    )
    tb(
        slide,
        12.05,
        1.48,
        7.3,
        0.45,
        [{"text": "ENGINEERING RISKS", "size": 22, "bold": True, "color": ORANGE, "after": 0}],
    )
    risks = [
        "Limited real failure labels. Synthetic trajectories are a development aid, not field proof.",
        "Synthetic-to-real gap: generated modes may not match defect physics.",
        "Corpus coverage is uneven. Several standards remain outside the retrieved set.",
        "Retrieval can miss or over-weight overlapping documents.",
        "An ungrounded hypothesis can still misstate a mechanism. Validation bounds this; it does not remove it.",
        "Sensor quality, model drift, and engineer trust are open deployment questions.",
        "No production validation is claimed.",
    ]
    ry = 2.15
    for item in risks:
        tb(slide, 11.25, ry, 8.25, 0.34, [{"text": "●   " + item, "size": 12, "color": INK, "after": 0}])
        ry += 0.35

    # 03 panel
    rrect(slide, 0.28, 4.92, 19.44, 3.55, RGBColor(0xEA, 0xF6, 0xFC), LINE, 0.06, 1.0)
    tb(
        slide,
        0.5,
        5.02,
        8.0,
        0.4,
        [
            {
                "runs": [
                    {"text": "03   ", "size": 18, "bold": True, "color": ROYAL},
                    {"text": "APPLICATION POTENTIAL", "size": 18, "bold": True, "color": BLUE},
                ],
                "after": 0,
            }
        ],
    )
    tb(
        slide,
        9.2,
        5.08,
        10.2,
        0.32,
        [{"text": "Where an investigation layer can sit beside existing reliability work", "size": 13, "color": MUTED, "align": PP_ALIGN.RIGHT}],
        align=PP_ALIGN.RIGHT,
    )

    # users
    rrect(slide, 0.48, 5.52, 6.3, 2.75, WHITE, radius=0.08)
    tb(slide, 0.6, 5.58, 6.0, 0.3, [{"text": "A    APPLICATION GROUPS", "size": 12, "bold": True, "color": BLUE, "after": 0}])
    users = [
        "Power-electronics engineering teams",
        "Reliability engineering teams",
        "EV and automotive electronics",
        "Industrial power electronics",
        "Module manufacturers and test laboratories",
        "R&D groups studying SiC degradation",
    ]
    uy = 5.95
    for u in users:
        tb(slide, 0.7, uy, 5.9, 0.32, [{"text": u, "size": 13, "color": INK, "after": 0}])
        uy += 0.35

    # value
    rrect(slide, 6.95, 5.52, 6.15, 2.75, WHITE, radius=0.08)
    tb(slide, 7.1, 5.58, 5.8, 0.3, [{"text": "B    VALUE CREATED", "size": 12, "bold": True, "color": BLUE, "after": 0}])
    values = [
        ("Faster investigation", "Less manual correlation across signals and documents."),
        ("Traceable evidence", "Provenance from observation to retrieved source."),
        ("Knowledge retention", "Investigations stay reproducible and reviewable."),
        ("Consistent reporting", "The same epistemic structure on every case."),
    ]
    vy = 5.98
    for name, note in values:
        tb(
            slide,
            7.15,
            vy,
            5.75,
            0.52,
            [
                {"text": name, "size": 13, "bold": True, "color": INK, "after": 0},
                {"text": note, "size": 11, "color": MUTED, "after": 0},
            ],
        )
        vy += 0.55

    # deployment
    rrect(slide, 13.28, 5.52, 6.2, 2.75, WHITE, radius=0.08)
    tb(slide, 13.42, 5.58, 5.9, 0.3, [{"text": "C    DEPLOYMENT MODEL", "size": 12, "bold": True, "color": BLUE, "after": 0}])
    deps = [
        (PALE, NAVY, "Existing tests and engineering data"),
        (RGBColor(0xE8, 0xEE, 0xFF), ROYAL, "+"),
        (RGBColor(0xE4, 0xF0, 0xFF), BLUE, "SmartESS analysis and investigation layer"),
        (RGBColor(0xEE, 0xE7, 0xFB), PURPLE, "↓"),
        (PALE_GREEN, GREEN, "Evidence-grounded report"),
        (RGBColor(0xE5, 0xF6, 0xEA), GREEN, "Engineer decision"),
    ]
    dy = 5.98
    for fill, color, text in deps:
        if text in {"+", "↓"}:
            tb(slide, 13.5, dy, 5.8, 0.22, [{"text": text, "size": 12, "bold": True, "color": color, "align": PP_ALIGN.CENTER, "after": 0}], align=PP_ALIGN.CENTER)
            dy += 0.2
        else:
            rrect(slide, 13.5, dy, 5.75, 0.32, fill, radius=0.1)
            tb(slide, 13.55, dy, 5.65, 0.32, [{"text": text, "size": 11, "bold": True, "color": color, "align": PP_ALIGN.CENTER, "after": 0}], align=PP_ALIGN.CENTER, anchor="ctr")
            dy += 0.36

    # 04
    rrect(slide, 0.28, 8.6, 19.44, 2.4, WHITE, LINE, 0.06, 1.0)
    tb(
        slide,
        0.48,
        8.7,
        10,
        0.35,
        [
            {
                "runs": [
                    {"text": "04   ", "size": 16, "bold": True, "color": ROYAL},
                    {"text": "VIABILITY", "size": 16, "bold": True, "color": BLUE},
                    {"text": "    A software intelligence layer. It does not replace hardware testing.", "size": 13, "color": MUTED},
                ],
                "after": 0,
            }
        ],
    )
    chips = [
        ("MODULE", "Investigate one module trajectory"),
        ("POPULATION", "Compare against the cohort"),
        ("TEST PROFILE", "Reuse across stress conditions"),
        ("LABORATORY", "Same workflow, many datasets"),
        ("INTEGRATION", "Versioned contracts and a REST API"),
        ("OPERATIONS", "Sits beside existing test practice"),
        ("MODELS", "Frozen, comparable analysis artifacts"),
        ("KNOWLEDGE", "Corpus grows by verified documents"),
    ]
    cx, cy = 0.48, 9.15
    for i, (name, note) in enumerate(chips):
        rrect(slide, cx, cy, 4.55, 0.78, PALE, radius=0.1)
        tb(
            slide,
            cx + 0.1,
            cy + 0.06,
            4.35,
            0.66,
            [
                {"text": name, "size": 12, "bold": True, "color": BLUE, "after": 0},
                {"text": note, "size": 12, "color": INK, "after": 0},
            ],
        )
        cx += 4.75
        if i == 3:
            cx = 0.48
            cy = 10.05


def slide_impact(slide):
    bg(slide, RGBColor(0xF3, 0xFB, 0xFE))
    team_badge(slide)
    slide_title(
        slide,
        "IMPACT & BENEFITS",
        "From a detected deviation to a reviewable engineering investigation",
        y=0.12,
    )
    agenticx_mark(slide)

    rrect(slide, 0.28, 1.2, 10.7, 6.55, RGBColor(0xE7, 0xF6, 0xFC), radius=0.08)
    tb(
        slide,
        0.45,
        1.32,
        7.4,
        0.55,
        [
            {"text": "Continuous engineering intelligence", "size": 16, "bold": True, "color": BLUE, "after": 0},
            {"text": "A closed investigation cycle. Learning stays inside reviewed reports.", "size": 12, "color": MUTED},
        ],
    )
    # center hub
    oval(slide, 4.15, 3.35, 3.0, 3.0, WHITE, ROYAL)
    oval(slide, 4.45, 3.65, 2.4, 2.4, RGBColor(0xE8, 0xF1, 0xFF))
    tb(
        slide,
        4.5,
        4.15,
        2.3,
        1.4,
        [
            {"text": "SmartESS", "size": 16, "bold": True, "color": NAVY, "align": PP_ALIGN.CENTER, "after": 0},
            {"text": "Engineering", "size": 12, "bold": True, "color": BLUE, "align": PP_ALIGN.CENTER, "after": 0},
            {"text": "Investigation", "size": 12, "bold": True, "color": BLUE, "align": PP_ALIGN.CENTER, "after": 0},
            {"text": "Cycle", "size": 12, "bold": True, "color": BLUE, "align": PP_ALIGN.CENTER, "after": 0},
        ],
        align=PP_ALIGN.CENTER,
        anchor="ctr",
    )
    nodes = [
        (4.85, 1.95, TEAL, "OBSERVE", "Module record"),
        (7.55, 2.55, ROYAL, "DETECT", "Anomaly evidence"),
        (8.35, 4.35, PURPLE, "INVESTIGATE", "Form the question"),
        (7.45, 6.15, ORANGE, "RETRIEVE", "Curated sources"),
        (4.85, 6.7, RED, "SYNTHESIZE", "Hypothesis + limits"),
        (2.15, 6.15, GREEN, "VALIDATE", "Support / contradict"),
        (1.25, 4.35, NAVY, "REPORT", "Structured case"),
        (2.15, 2.55, RGBColor(0x0E, 0x74, 0x86), "CAPTURE", "Retain the case"),
    ]
    for x, y, color, name, note in nodes:
        oval(slide, x, y, 1.55, 0.85, color)
        tb(
            slide,
            x,
            y + 0.08,
            1.55,
            0.7,
            [
                {"text": name, "size": 11, "bold": True, "color": WHITE, "align": PP_ALIGN.CENTER, "after": 0},
                {"text": note, "size": 9, "color": RGBColor(0xF4, 0xF8, 0xFF), "align": PP_ALIGN.CENTER},
            ],
            align=PP_ALIGN.CENTER,
            anchor="ctr",
        )

    # right benefits
    rrect(slide, 11.15, 1.2, 8.55, 6.55, WHITE, LINE, 0.08, 1.0)
    rect(slide, 11.15, 1.2, 8.55, 0.7, ROYAL)
    tb(
        slide,
        11.3,
        1.28,
        5.2,
        0.55,
        [{"text": "Key Benefits", "size": 20, "bold": True, "color": WHITE, "after": 0}],
        anchor="ctr",
    )
    tb(
        slide,
        15.3,
        1.32,
        4.15,
        0.48,
        [{"text": "Categories of value.\nNo savings figures are claimed.", "size": 11, "color": RGBColor(0xD6, 0xE6, 0xFF), "align": PP_ALIGN.RIGHT}],
        align=PP_ALIGN.RIGHT,
        anchor="ctr",
    )
    benefits = [
        (0, 0, PALE_GREEN, GREEN, "ECONOMIC", ["Less manual investigation effort", "Faster path from signal to a reviewable case"]),
        (1, 0, PALE, ROYAL, "INDUSTRIAL", ["A repeatable investigation workflow", "Clearer reliability-engineering visibility"]),
        (0, 1, PALE_PURPLE, PURPLE, "KNOWLEDGE", ["Structured retention of cases", "Evidence and provenance preserved"]),
        (1, 1, PALE_ORANGE, ORANGE, "OPERATIONAL", ["Faster move from anomaly to investigation", "Consistent technical reporting"]),
    ]
    for col, row, fill, color, title, lines in benefits:
        x = 11.4 + col * 4.05
        y = 2.15 + row * 2.65
        rrect(slide, x, y, 3.9, 2.45, fill, radius=0.1)
        tb(slide, x + 0.15, y + 0.15, 3.6, 0.4, [{"text": title, "size": 16, "bold": True, "color": color, "after": 0}])
        ly = y + 0.7
        for line in lines:
            tb(slide, x + 0.18, ly, 3.55, 0.7, [{"text": "•  " + line, "size": 13, "color": INK, "after": 0}])
            ly += 0.7

    # bottom comparison
    rect(slide, 0, 7.9, SW, 3.35, WHITE)
    tb(
        slide,
        0.4,
        8.0,
        19.2,
        0.35,
        [{"text": "FROM FRAGMENTED REVIEW TO A TRACEABLE INVESTIGATION", "size": 14, "bold": True, "color": NAVY, "align": PP_ALIGN.CENTER, "after": 0}],
        align=PP_ALIGN.CENTER,
    )
    rrect(slide, 0.35, 8.45, 8.3, 2.5, PALE_RED, radius=0.08)
    tb(
        slide,
        0.5,
        8.55,
        8.0,
        0.35,
        [{"text": "CURRENT STATE", "size": 14, "bold": True, "color": RED, "after": 0}],
    )
    current = [
        "Manual correlation of signals, limits, and documents",
        "In-spec abnormal behavior often has no investigation record",
        "Reasoning is difficult to reproduce or audit",
        "Literature and measurements are easy to mix together",
    ]
    cy = 9.0
    for c in current:
        tb(slide, 0.6, cy, 7.8, 0.4, [{"text": "•  " + c, "size": 13, "color": INK, "after": 0}])
        cy += 0.42

    chevron(slide, 8.85, 9.2, 1.9, 0.7, ROYAL)
    tb(
        slide,
        8.9,
        9.32,
        1.7,
        0.45,
        [{"text": "SmartESS", "size": 12, "bold": True, "color": WHITE, "align": PP_ALIGN.CENTER, "after": 0}],
        align=PP_ALIGN.CENTER,
    )

    rrect(slide, 11.0, 8.45, 8.65, 2.5, PALE_GREEN, radius=0.08)
    tb(
        slide,
        11.15,
        8.55,
        8.3,
        0.35,
        [{"text": "WITH SMARTESS", "size": 14, "bold": True, "color": GREEN, "after": 0}],
    )
    future = [
        "Agentic, evidence-grounded engineering investigation",
        "Measured data, model output, sources, and hypotheses stay separate",
        "Every important retrieved claim points at its source",
        "The engineer decides. The system does not release hardware.",
    ]
    fy = 9.0
    for c in future:
        tb(slide, 11.25, fy, 8.2, 0.4, [{"text": "•  " + c, "size": 13, "color": INK, "after": 0}])
        fy += 0.42


def slide_references(slide):
    bg(slide, WHITE)
    team_badge(slide)
    slide_title(
        slide,
        "RESEARCH & REFERENCES",
        "Verified corpus sources, plus the methods this prototype actually runs",
        y=0.12,
    )
    agenticx_mark(slide)

    tb(
        slide,
        0.4,
        1.12,
        19.2,
        0.32,
        [
            {
                "runs": [
                    {"text": "Key references   ", "size": 16, "bold": True, "color": BLUE},
                    {
                        "text": "These six cards are inspected corpus documents. Detection, retrieval, and the language model are stated in the band below.",
                        "size": 12,
                        "color": MUTED,
                    },
                ],
                "after": 0,
            }
        ],
    )

    cards = [
        {
            "role": "STANDARD  ·  CORPUS",
            "org": "ECPE",
            "title": "AQG 324, Release 04.1/2025",
            "body": "Qualification of power modules for automotive converters, including the SiC annex: power cycling, thermal cycling, HTRB, HTGB, and thermal resistance.",
            "why": "Defines test programmes. It does not map one measured parameter to one failure mechanism.",
            "url": "https://www.ecpe.org/research/working-groups/automotive-aqg-324/",
        },
        {
            "role": "SiC DEGRADATION  ·  CORPUS",
            "org": "Energies, 2024",
            "title": "Forward power cycling of discrete SiC MOSFETs",
            "body": "Huang, Singh, Liu, and Norrga. Failure characterization under forward power cycling.",
            "why": "Multi-parameter evidence: RDS(on), Vth, leakage, and thermal quantities together.",
            "url": "https://www.mdpi.com/1996-1073/17/11/2557",
        },
        {
            "role": "ELECTRICAL / THERMAL  ·  CORPUS",
            "org": "Wolfspeed, 2025",
            "title": "Power cycling and lifetime modeling",
            "body": "Ceresa and Shaw. SiC power-module power cycling, package mechanisms, and lifetime models.",
            "why": "Thermal resistance and junction temperature are monitored inputs, not a single-cause diagnosis.",
            "url": "https://assets.wolfspeed.com/uploads/2025/11/Wolfspeed_Power_Cycling_and_Lifetime_Modeling_Approach.pdf",
        },
        {
            "role": "PACKAGING RELIABILITY  ·  CORPUS",
            "org": "Energies, 2022",
            "title": "Wide-bandgap semiconductor and packaging review",
            "body": "Wang, Ding, and Yin. Gate oxide, bond wire, die attach, and thermal-path mechanisms.",
            "why": "Distinct mechanisms can share the same electrical and thermal observables.",
            "url": "https://www.mdpi.com/1996-1073/15/18/6670",
        },
        {
            "role": "GATE OXIDE  ·  CORPUS",
            "org": "NIST",
            "title": "SiC MOSFET gate-oxide breakdown status",
            "body": "K. P. Cheung, NIST Engineering Physics Division. Intrinsic versus extrinsic oxide failure.",
            "why": "Argues against simplistic reliability extrapolation. Supports bounded claims.",
            "url": "https://tsapps.nist.gov/publication/get_pdf.cfm?pub_id=925228",
        },
        {
            "role": "PROGNOSIS REVIEW  ·  CORPUS",
            "org": "Entropy, 2026",
            "title": "Prognosis approaches for power SiC MOSFETs",
            "body": "Kumar et al. Health-state and remaining-useful-life indicators under stress.",
            "why": "Indicator-to-mechanism links are ambiguous. SmartESS does not claim remaining-life prediction.",
            "url": "https://www.mdpi.com/1099-4300/28/2/234",
        },
    ]
    cw = 3.15
    for i, card in enumerate(cards):
        x = 0.28 + i * (cw + 0.08)
        rrect(slide, x, 1.52, cw, 5.15, PALE, radius=0.1)
        tb(slide, x + 0.08, 1.6, cw - 0.16, 0.55, [{"text": card["role"], "size": 10, "bold": True, "color": TEAL, "align": PP_ALIGN.CENTER, "after": 0}], align=PP_ALIGN.CENTER)
        tb(slide, x + 0.1, 2.15, cw - 0.2, 0.7, [{"text": card["org"], "size": 14, "bold": True, "color": NAVY, "align": PP_ALIGN.CENTER, "after": 0}], align=PP_ALIGN.CENTER)
        tb(slide, x + 0.1, 2.8, cw - 0.2, 0.85, [{"text": card["title"], "size": 13, "bold": True, "color": BLUE, "align": PP_ALIGN.CENTER, "after": 0}], align=PP_ALIGN.CENTER)
        tb(slide, x + 0.1, 3.65, cw - 0.2, 1.15, [{"text": card["body"], "size": 11, "color": INK, "align": PP_ALIGN.CENTER, "after": 0}], align=PP_ALIGN.CENTER)
        tb(slide, x + 0.1, 4.8, cw - 0.2, 1.05, [{"text": card["why"], "size": 11, "italic": True, "color": MUTED, "align": PP_ALIGN.CENTER, "after": 0}], align=PP_ALIGN.CENTER)
        rrect(slide, x + 0.35, 5.95, cw - 0.7, 0.48, ROYAL, radius=0.2)
        tb(
            slide,
            x + 0.35,
            5.98,
            cw - 0.7,
            0.42,
            [{"text": "VIEW SOURCE", "size": 11, "bold": True, "color": WHITE, "align": PP_ALIGN.CENTER, "url": card["url"], "after": 0}],
            align=PP_ALIGN.CENTER,
            anchor="ctr",
        )

    # bottom: methods + knowledge pipeline + UI surfaces (no borrowed screenshots)
    rect(slide, 0, 6.85, SW, 4.4, RGBColor(0x0E, 0x1B, 0x2E))
    tb(
        slide,
        0.4,
        6.95,
        19.2,
        0.32,
        [{"text": "KNOWLEDGE PIPELINE   ·   METHODS IN CODE   ·   ENGINEER INTERFACE", "size": 13, "bold": True, "color": WHITE, "after": 0}],
    )
    steps = [
        "Engineering\ndocuments",
        "Extraction",
        "Chunking",
        "Metadata &\nprovenance",
        "Embeddings\nMiniLM",
        "ChromaDB",
        "Retriever",
        "Investigation\nagents",
    ]
    sx = 0.35
    for i, step in enumerate(steps):
        rrect(slide, sx, 7.4, 2.15, 0.85, RGBColor(0x1A, 0x3A, 0x66), radius=0.1)
        tb(
            slide,
            sx,
            7.45,
            2.15,
            0.75,
            [{"text": step, "size": 11, "bold": True, "color": WHITE, "align": PP_ALIGN.CENTER, "after": 0}],
            align=PP_ALIGN.CENTER,
            anchor="ctr",
        )
        if i < len(steps) - 1:
            tb(slide, sx + 1.95, 7.55, 0.35, 0.4, [{"text": "→", "size": 14, "bold": True, "color": RGBColor(0x8E, 0xB6, 0xFF), "after": 0}])
        sx += 2.45

    shots = [
        ("mission.png", "Mission Control", "Stage chain and readiness"),
        ("evidence.png", "Evidence", "Retrieved passages with provenance"),
        ("report.png", "Report", "Observed, calculated, hypothesized"),
    ]
    root = Path(__file__).resolve().parents[1] / "docs" / "deck-assets"
    sx = 0.28
    for filename, title, note in shots:
        path = root / filename
        rrect(slide, sx, 8.38, 4.35, 2.72, RGBColor(0x15, 0x2C, 0x4E), radius=0.06)
        tb(
            slide,
            sx + 0.08,
            8.4,
            4.2,
            0.28,
            [{"text": f"{title}  ·  {note}", "size": 10, "bold": True, "color": RGBColor(0xC5, 0xDC, 0xFF), "after": 0}],
        )
        if path.exists():
            slide.shapes.add_picture(str(path), Inches(sx + 0.08), Inches(8.7), Inches(4.19), Inches(2.32))
        sx += 4.48
    rrect(slide, 13.75, 8.38, 5.95, 2.72, RGBColor(0x15, 0x2C, 0x4E), radius=0.08)
    tb(
        slide,
        13.9,
        8.48,
        5.65,
        2.5,
        [
            {"text": "METHODS IN THIS PROTOTYPE", "size": 12, "bold": True, "color": RGBColor(0x9D, 0xC2, 0xFF), "after": 4},
            {"text": "Anomaly score: Isolation Forest (Liu, Ting, Zhou, ICDM 2008) via scikit-learn. A score is a model output.", "size": 12, "color": WHITE, "after": 6},
            {"text": "Retrieval: curated RAG over the inspected corpus only. Lewis et al., NeurIPS 2020, is the method foundation — not an open-web search.", "size": 12, "color": WHITE, "after": 6},
            {"text": "Hypotheses: Llama 3.3 70B via OpenRouter, from supplied evidence. The model does not measure the module.", "size": 12, "color": WHITE, "after": 0},
        ],
    )


if __name__ == "__main__":
    out = "docs/AgenticX-SmartESS.pptx"
    prs = build()
    prs.save(out)
    print(out, "slides", len(prs.slides))
